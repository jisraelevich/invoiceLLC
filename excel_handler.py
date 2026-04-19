"""
Módulo de gestión de archivos Excel.
Encargado de leer, filtrar y actualizar el archivo de facturas.
"""

from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from datetime import datetime
from typing import List, Dict

from utils import LoggerFactory

logger = LoggerFactory.get_logger(__name__)


class ExcelHandler:
    """Gestiona la lectura y actualización del archivo Excel de facturas."""
    
    def __init__(self, archivo_excel: str = "facturas.xlsx"):
        """
        Inicializa el manejador de Excel.
        
        Args:
            archivo_excel: Ruta del archivo Excel
        """
        self.archivo_excel = Path(archivo_excel)
        self.workbook = None
        self.worksheet = None
        
    def cargar_workbook(self) -> bool:
        """
        Carga el workbook de Excel.
        
        Returns:
            bool: True si se cargó exitosamente, False sino
        """
        try:
            if not self.archivo_excel.exists():
                logger.error(f"Archivo {self.archivo_excel} no encontrado")
                return False
            
            self.workbook = load_workbook(self.archivo_excel)
            self.worksheet = self.workbook.active
            logger.info(f"Workbook cargado: {self.archivo_excel}")
            return True
        except Exception as e:
            logger.error(f"Error al cargar workbook: {e}")
            return False
    
    def obtener_filas_pendientes(self) -> List[Dict]:
        """
        Obtiene todas las filas con estado = PENDIENTE.
        
        Returns:
            List[Dict]: Lista de diccionarios con los datos de cada fila
        """
        if self.worksheet is None:
            logger.error("Workbook no cargado")
            return []
        
        filas_pendientes = []
        
        # Obtener headers de la primera fila
        headers = []
        for cell in self.worksheet[1]:
            headers.append(cell.value)
        
        if not headers or 'estado' not in [h.lower() if h else '' for h in headers]:
            logger.error("Headers no encontrados o 'estado' no existe")
            return []
        
        # Iterar desde la fila 2 (la 1 es header)
        for idx, row in enumerate(self.worksheet.iter_rows(min_row=2, values_only=False), start=2):
            fila_dict = {}
            
            # Mapear columnas a valores
            for col_idx, (header, cell) in enumerate(zip(headers, row)):
                fila_dict[header.lower() if header else f'col_{col_idx}'] = cell.value
            
            # Solo procesar filas con estado PENDIENTE
            if fila_dict.get('estado', '').upper() == 'PENDIENTE':
                fila_dict['numero_fila'] = idx  # Guardar el número de fila para actualizar después
                filas_pendientes.append(fila_dict)
        
        logger.info(f"Se encontraron {len(filas_pendientes)} facturas pendientes")
        return filas_pendientes
    
    def actualizar_fila(self, numero_fila: int, cae: str, nro_comprobante: str) -> bool:
        """
        Actualiza una fila con el CAE y número de comprobante.
        
        Args:
            numero_fila: Número de fila a actualizar (1-indexed)
            cae: Código de Autorización Electrónica
            nro_comprobante: Número de comprobante
            
        Returns:
            bool: True si se actualizó exitosamente
        """
        try:
            if self.worksheet is None:
                logger.error("Workbook no cargado")
                return False
            
            # Obtener headers para encontrar las columnas
            headers = [cell.value for cell in self.worksheet[1]]
            
            cae_col = None
            estado_col = None
            nro_col = None
            
            for idx, header in enumerate(headers, start=1):
                if header and header.lower() == 'cae':
                    cae_col = idx
                elif header and header.lower() == 'estado':
                    estado_col = idx
                elif header and header.lower() == 'nro_comprobante':
                    nro_col = idx
            
            # Actualizar células
            if cae_col:
                self.worksheet.cell(row=numero_fila, column=cae_col).value = cae
            if nro_col:
                self.worksheet.cell(row=numero_fila, column=nro_col).value = nro_comprobante
            if estado_col:
                self.worksheet.cell(row=numero_fila, column=estado_col).value = "EMITIDA"
                # Marcar como verde para indicar que fue emitida
                green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
                for col_idx in range(1, len(headers) + 1):
                    self.worksheet.cell(row=numero_fila, column=col_idx).fill = green_fill
            
            logger.info(f"Fila {numero_fila} actualizada: CAE={cae}, Nro={nro_comprobante}")
            return True
        except Exception as e:
            logger.error(f"Error al actualizar fila {numero_fila}: {e}")
            return False
    
    def guardar(self) -> bool:
        """
        Guarda los cambios en el archivo Excel.
        
        Returns:
            bool: True si se guardó exitosamente
        """
        try:
            if self.workbook is None:
                logger.error("Workbook no cargado")
                return False
            
            self.workbook.save(self.archivo_excel)
            logger.info(f"Workbook guardado: {self.archivo_excel}")
            return True
        except Exception as e:
            logger.error(f"Error al guardar workbook: {e}")
            return False
    
    def cerrar(self):
        """Cierra el workbook."""
        if self.workbook:
            self.workbook.close()


def crear_excel_ejemplo(archivo: str = "facturas.xlsx") -> bool:
    """
    Crea un archivo Excel de ejemplo para pruebas.
    
    Args:
        archivo: Nombre del archivo a crear
        
    Returns:
        bool: True si se creó exitosamente
    """
    try:
        from openpyxl import Workbook
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Facturas"
        
        # Headers
        headers = ["fecha", "descripcion", "monto", "cuit_cliente", "nombre_cliente", "estado", "cae", "nro_comprobante"]
        ws.append(headers)
        
        # Fila de ejemplo 1
        ws.append([
            "02/04/2026",
            "Servicio de consultoría",
            "5000",
            "99999999999",
            "Consumidor Final",
            "PENDIENTE",
            "",
            ""
        ])
        
        # Fila de ejemplo 2
        ws.append([
            "02/04/2026",
            "Soporte técnico mensual",
            "2500",
            "20123456789",
            "Empresa XYZ SRL",
            "PENDIENTE",
            "",
            ""
        ])
        
        wb.save(archivo)
        logger.info(f"Excel de ejemplo creado: {archivo}")
        return True
    except Exception as e:
        logger.error(f"Error al crear Excel de ejemplo: {e}")
        return False
