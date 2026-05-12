import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

# Rutas
JSON_PATH = Path("..") / "data" / "2026" / "april_2026.json"
REPORTE_XLSX = Path("..") / "Facturas" / "reporte_facturacion_abril_2026.xlsx"

# Cargar datos
with open(JSON_PATH, encoding="utf-8") as f:
    data = json.load(f)

facturas = data.get("facturas", [])

# Crear workbook y hoja
wb = Workbook()
ws = wb.active
ws.title = "Resumen Facturación"

# Encabezados
headers = [
    "Cliente", "CUIT", "Concepto", "Fecha Planificada", "Monto", "Tipo", "Estado", "CAE", "Nro Comprobante", "Fecha Emisión"
]
ws.append(headers)
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
for col in range(1, len(headers)+1):
    ws.cell(row=1, column=col).font = header_font
    ws.cell(row=1, column=col).fill = header_fill

# Agregar datos
for f in facturas:
    ws.append([
        f.get("cliente", ""),
        f.get("cuit", ""),
        f.get("concepto", ""),
        f.get("fecha_planificada", ""),
        f.get("monto", ""),
        f.get("tipo", ""),
        f.get("estado", ""),
        f.get("cae", ""),
        f.get("nro_comprobante", ""),
        f.get("fecha_emision", "")
    ])

# Autosize columns
for col in ws.columns:
    max_length = 0
    col_letter = col[0].column_letter
    for cell in col:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
    ws.column_dimensions[col_letter].width = max_length + 2

wb.save(REPORTE_XLSX)
print(f"[OK] Reporte generado: {REPORTE_XLSX}")
